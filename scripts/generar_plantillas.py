"""Genera plantillas/Plantilla_Ficha_SUM_<COMP>.xlsx desde el diccionario (una por componente).

Uso:  python scripts/generar_plantillas.py
"""
import csv, json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

RAIZ = pathlib.Path(__file__).resolve().parent.parent
DIC = list(csv.DictReader(open(RAIZ / "estandar" / "diccionario_campos.csv", encoding="utf-8")))
LISTAS = json.load(open(RAIZ / "estandar" / "listas.json", encoding="utf-8"))
OUT = RAIZ / "plantillas"

COLS = ["Sección", "Clave", "Campo", "Tipo", "Obligatorio", "Estado", "Valor", "Mín", "Máx", "Calificador",
        "Unidad", "Fuentes (ID; ID)", "Localizador", "Evidencia", "Certeza", "Justificación (escalas)", "Nota"]
ANCHO = [16, 22, 34, 12, 8, 13, 40, 9, 9, 10, 13, 16, 14, 9, 16, 30, 40]
AZUL, GRIS, AMAR = "1F3A5F", "EDEFF2", "FFF6D5"
fino = Side(style="thin", color="C9CDD3")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)

TIPO_TXT = {"numero": "número", "entero": "entero", "escala": "escala 1–5", "lista": "lista", "lista_multiple": "lista múltiple (|)",
            "booleano": "Sí/No/Cond.", "valor_o_clase": "valor o clase", "texto": "texto", "texto_largo": "texto largo"}


def dom(n):
    v = LISTAS[n]
    return list(v.keys()) if isinstance(v, dict) else v


def hoja_listas(wb):
    ws = wb.create_sheet("Listas")
    nombres = ["estado", "calificador", "evidencia", "certeza", "tipo_fuente"] + sorted(
        {r["dominio"] for r in DIC if r["dominio"]})
    for j, n in enumerate(nombres, start=1):
        vals = dom(n)
        ws.cell(1, j, n).font = Font(bold=True)
        for i, v in enumerate(vals, start=2):
            ws.cell(i, j, v)
        col = get_column_letter(j)
        wb.defined_names[f"L_{n}"] = DefinedName(f"L_{n}", attr_text=f"Listas!${col}$2:${col}${len(vals) + 1}")
        ws.column_dimensions[col].width = 18
    # unidades permitidas por campo
    j = len(nombres) + 2
    ws.cell(1, j, "unidades por campo").font = Font(bold=True)
    return ws


def generar(comp):
    info = LISTAS["component"][comp]
    filas = [r for r in DIC if r["componentes"] == "TODOS" or comp in r["componentes"].split("|")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Ficha"
    ws["A1"] = f"FICHA SUM v1.0 · {info['nombre'].upper()}"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"], ws["C2"], ws["E2"], ws["G2"] = "Autor", "", "ID ficha", f"{comp}-XXX-01"
    ws["A3"] = ("Una plantilla = una ficha. Estado: dato / no_disponible / no_aplica / pendiente. "
                "Números en 'Valor' (o 'Mín'–'Máx' si es rango) y unidad en su columna. Toda fuente se registra en la hoja Fuentes y aquí se cita por ID.")
    ws.merge_cells("A3:Q3")
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 32
    for c in ("A2", "E2"):
        ws[c].font = Font(bold=True)
    for j, (h, w) in enumerate(zip(COLS, ANCHO), start=1):
        cel = ws.cell(5, j, h)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=AZUL)
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "G6"

    dv = {n: DataValidation(type="list", formula1=f"=L_{n}", allow_blank=True) for n in ["estado", "calificador", "evidencia", "certeza"]}
    dv_esc = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                            error="Escala entera 1–5", showErrorMessage=True)
    dv_num = DataValidation(type="decimal", operator="between", formula1="-1E12", formula2="1E12", allow_blank=True,
                            error="Solo números: el rango va en Mín/Máx y la unidad en su columna", showErrorMessage=True)
    dvs_dom, dvs_uni = {}, {}
    for v in list(dv.values()) + [dv_esc, dv_num]:
        ws.add_data_validation(v)

    seccion_prev = None
    for i, r in enumerate(filas, start=6):
        vals = [r["seccion"] if r["seccion"] != seccion_prev else "", r["clave"], r["campo"], TIPO_TXT[r["tipo"]],
                r["obligatorio"], "pendiente", "", "", "", "", r["unidad"] if r["unidad"] not in ("1–5",) else "",
                "", "", "", "", "", r["regla"]]
        seccion_prev = r["seccion"]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(i, j, v)
            c.border = BORDE
            c.alignment = Alignment(wrap_text=j in (3, 7, 16, 17), vertical="top")
            if j <= 5:
                c.fill = PatternFill("solid", fgColor=GRIS)
        ws.cell(i, 1).font = Font(bold=True, color=AZUL)
        ws.cell(i, 2).font = Font(name="Consolas", size=9)
        for col, n in ((6, "estado"), (10, "calificador"), (14, "evidencia"), (15, "certeza")):
            dv[n].add(ws.cell(i, col))
        if r["tipo"] == "escala":
            dv_esc.add(ws.cell(i, 7))
        elif r["tipo"] in ("numero", "entero"):
            dv_num.add(ws.cell(i, 7)); dv_num.add(ws.cell(i, 8)); dv_num.add(ws.cell(i, 9))
        elif r["tipo"] in ("lista", "booleano") and r["dominio"]:
            if r["dominio"] not in dvs_dom:
                dvs_dom[r["dominio"]] = DataValidation(type="list", formula1=f"=L_{r['dominio']}", allow_blank=True)
                ws.add_data_validation(dvs_dom[r["dominio"]])
            dvs_dom[r["dominio"]].add(ws.cell(i, 7))
        unis = [u for u in [r["unidad"]] + r["unidades_alternativas"].split("|") if u and u != "1–5"]
        if len(unis) > 1:
            key = "|".join(unis)
            if key not in dvs_uni:
                dvs_uni[key] = DataValidation(type="list", formula1='"' + ",".join(unis) + '"', allow_blank=True)
                ws.add_data_validation(dvs_uni[key])
            dvs_uni[key].add(ws.cell(i, 11))
    last = 5 + len(filas)
    # resaltado: obligatorio aún pendiente / dato sin evidencia
    ws.conditional_formatting.add(f"F6:F{last}", FormulaRule(formula=[f'AND($E6="Sí",$F6="pendiente")'],
                                                             fill=PatternFill("solid", fgColor="F8D7DA")))
    ws.conditional_formatting.add(f"N6:N{last}", FormulaRule(formula=[f'AND($F6="dato",$N6="")'],
                                                             fill=PatternFill("solid", fgColor=AMAR)))
    ws.auto_filter.ref = f"A5:Q{last}"

    wf = wb.create_sheet("Fuentes")
    for j, (h, w) in enumerate(zip(["ID", "Cita (APA 7)", "URL", "Tipo", "Fecha de consulta"], [10, 80, 40, 20, 16]), start=1):
        c = wf.cell(1, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=AZUL)
        wf.column_dimensions[get_column_letter(j)].width = w
    wf.cell(2, 1, "PROPIO"); wf.cell(2, 2, "Elaboración propia del estudiante."); wf.cell(2, 4, "elaboracion_propia")
    dvt = DataValidation(type="list", formula1="=L_tipo_fuente", allow_blank=True); wf.add_data_validation(dvt); dvt.add("D2:D200")

    wi = wb.create_sheet("Instrucciones", 0)
    txt = [
        ("FICHA SUM v1.0 — cómo llenar", True),
        ("1. Completá Autor e ID en la hoja Ficha (formato " + comp + "-SUB-NN, p. ej. " + comp + "-ABC-01).", False),
        ("2. Para cada campo elegí el Estado. 'pendiente' en un campo obligatorio se marca en rojo.", False),
        ("3. Número puntual → Valor. Rango → Mín y Máx (Calificador = rango). Aproximado → Calificador ≈.", False),
        ("4. Nunca escribas unidades ni comentarios dentro de Valor: van en Unidad y Nota.", False),
        ("5. Listas múltiples: separá con |  (ej. Vivienda | Salud).", False),
        ("6. Fuentes: registrá cada referencia una vez en la hoja Fuentes y citála por ID (ej. F01; F03). Página o tabla → Localizador.", False),
        ("7. Evidencia (tipo de fuente): E1 criterio propio · E2 fabricante · E3 norma/ensayo · E4 independiente · E5 varias.", False),
        ("8. Certeza (calidad del número): verificado · rango_literatura · estimado_analogia. E1 nunca es 'verificado'.", False),
        ("9. Escalas 1–5: la Justificación es obligatoria (qué criterio sostiene el puntaje).", False),
        ("10. Valor de laboratorio ≠ desempeño en servicio en clima tropical húmedo: dejalo explícito en la Nota.", False),
        ("11. Convertí a JSON con:  python scripts/importar_plantilla.py <archivo.xlsx>  y validá con  python scripts/validar.py", False),
    ]
    for i, (t, b) in enumerate(txt, start=1):
        wi.cell(i, 1, t).font = Font(bold=b, size=13 if b else 11, color=AZUL if b else "000000")
    wi.column_dimensions["A"].width = 130
    hoja_listas(wb)
    OUT.mkdir(exist_ok=True)
    p = OUT / f"Plantilla_Ficha_SUM_{comp}.xlsx"
    wb.active = 1
    wb.save(p)
    return p, len(filas)


if __name__ == "__main__":
    for comp in LISTAS["component"]:
        p, n = generar(comp)
        print(f"{p.name}: {n} campos")
