"""Migra fichas TP1 en formato xlsx legado (plantilla 'FICHA ESTANDARIZADA', columnas A–J)
al estándar Ficha SUM v1.0: un archivo JSON por solución en data/<componente>/<ID>.json.

Uso:
    python scripts/migrar_xlsx.py <carpeta_con_xlsx> [--mapa estandar/mapa_ids_migracion.csv]

Qué hace:
  * Lee cada pestaña que tenga 'Clave APP' en B4 (ignora plantillas vacías).
  * Renombra claves legadas a las del diccionario (columna clave_legacy).
  * Separa número, rango, calificador (≈ ≥ ≤) y unidad; lo que no logra convertir lo deja en
    'valor_texto_original' para que el estudiante lo resuelva (el validador lo marca como error).
  * Convierte las fuentes en una bibliografía con ID y referencias por campo.
  * NUNCA inventa certeza: ese atributo queda vacío y lo completa el estudiante.
"""
import csv, json, re, sys, pathlib, datetime, unicodedata, argparse
import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
EST = RAIZ / "estandar"
DIC = {r["clave"]: r for r in csv.DictReader(open(EST / "diccionario_campos.csv", encoding="utf-8"))}
LISTAS = json.load(open(EST / "listas.json", encoding="utf-8"))
LEGADO = {}
for k, r in DIC.items():
    for old in filter(None, r["clave_legacy"].split("|")):
        LEGADO.setdefault(old, []).append(k)
CARPETA_COMP = {c: v["nombre"].split(" ")[0].lower() for c, v in LISTAS["component"].items()}

# ---------------------------------------------------------------- utilidades

def sin_tildes(s):
    return "".join(ch for ch in unicodedata.normalize("NFD", s) if unicodedata.category(ch) != "Mn").lower().strip()


def dominio(nombre):
    if not nombre:
        return None
    v = LISTAS[nombre]
    return list(v.keys()) if isinstance(v, dict) else v


def en_dominio(texto, dom):
    """Devuelve el valor canónico del dominio si coincide (sin tildes ni mayúsculas)."""
    t = sin_tildes(texto)
    for d in dom:
        if sin_tildes(d) == t:
            return d
    return None


NUM = r"[-+]?\d{1,3}(?:[   ]\d{3})+(?:[.,]\d+)?|[-+]?\d+(?:[.,]\d+)?"


def a_float(s):
    s = s.replace(" ", " ").replace(" ", " ").strip()
    if re.fullmatch(r"[-+]?\d{1,3}(?: \d{3})+(?:[.,]\d+)?", s):
        s = s.replace(" ", "")
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    elif s.count(",") >= 1 and s.count(".") == 1:
        s = s.replace(",", "")
    return float(s)


def parse_numero(v):
    """Interpreta números legados. Devuelve dict parcial o None si no es convertible sin pérdida."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return {"valor": round(float(v), 6) if not float(v).is_integer() else int(v), "calificador": "="}
    s = str(v).strip()
    s = re.sub(r"\s*\([^()]*\)\s*$", "", s)                       # comentario final entre paréntesis
    s = re.sub(r"\s*(mm|cm|m|años|%|kg/m²|kg/m2|kg/m³|kg/m3|MPa|kPa|Pa|dB|min|meses|km|N)\.?$", "", s)
    s2 = re.sub(r"^(₡|CRC|USD|US\$|\$)\s*", "", s)
    m = re.fullmatch(rf"(≈|~|±|≥|>=|>|≤|<=|<)?\s*({NUM})\s*(\+)?", s2)
    if m:
        cal = {"≈": "≈", "~": "≈", "≥": "≥", ">=": "≥", ">": "≥", "≤": "≤", "<=": "≤", "<": "≤"}.get(m.group(1) or "", "=")
        if m.group(3):
            cal = "≥"
        return {"valor": a_float(m.group(2)), "calificador": cal}
    m = re.fullmatch(rf"(≈|~)?\s*(?:₡|CRC|USD|\$)?\s*({NUM})\s*(?:–|-|—|a|to)\s*(?:₡|CRC|USD|\$)?\s*({NUM})", s2)
    if m:
        a, b = a_float(m.group(2)), a_float(m.group(3))
        if a <= b:
            return {"min": a, "max": b, "calificador": "rango"}
    return None


UNIDAD_ALIAS = {
    "W/m·K": "W/(m·K)", "W/mK": "W/(m·K)", "W/m²·K": "W/(m²·K)", "W/m2K": "W/(m²·K)",
    "kgCO₂e/m²": "kgCO2e/m²", "kgCO₂e/kg": "kgCO2e/kg", "kgCO₂e/unidad": "kgCO2e/unidad", "kgCO₂e/m³": "kgCO2e/m³",
    "₡/m²": "CRC/m²", "Rw (dB)": "dB (Rw)", "dB (LIA, lluvia intensa)": "dB (LIA)",
}


def unidad_canonica(clave, legado):
    d = DIC[clave]
    permitidas = [u for u in [d["unidad"]] + d["unidades_alternativas"].split("|") if u]
    u = UNIDAD_ALIAS.get((legado or "").strip(), (legado or "").strip())
    if u in permitidas:
        return u
    return d["unidad"] or None


# ---------------------------------------------------------------- fuentes

PROPIAS = ("elaboración propia", "elaboracion propia", "asignado por el estudiante", "investigación propia", "criterio propio")
URL = re.compile(r"https?://\S+")


class Bibliografia:
    def __init__(self):
        self.items = {}      # id -> dict
        self.por_texto = {}  # texto -> id
        self.n = 0

    def registrar(self, id_, cita, tipo=None):
        cita = re.sub(r"\s+", " ", str(cita)).strip()
        e = {"id": id_, "cita": cita}
        u = URL.search(cita)
        if u:
            e["url"] = u.group(0).rstrip(".,;)")
        if tipo:
            e["tipo"] = tipo
        self.items[id_] = e
        self.por_texto[cita] = id_
        return id_

    def ref_texto(self, texto):
        t = re.sub(r"\s+", " ", texto).strip()
        if t in self.por_texto:
            return self.por_texto[t]
        self.n += 1
        return self.registrar(f"R{self.n:02d}", t)

    def resolver(self, celda):
        """Convierte el contenido de la columna Fuente en una lista de referencias."""
        if celda is None:
            return []
        s = str(celda).strip()
        if s in ("", "—", "-", "–") or sin_tildes(s) in ("no disponible", "no aplica", "n/a"):
            return []
        if sin_tildes(s).startswith(PROPIAS):
            if "PROPIO" not in self.items:
                self.registrar("PROPIO", "Elaboración propia del estudiante (TP1).", "elaboracion_propia")
            return [{"ref": "PROPIO"}]
        partes = [p.strip() for p in re.split(r"\s*[;·]\s*|\s*,\s*(?=[A-Z]{1,3}\d)", s) if p.strip()]
        refs = []
        for p in partes:
            if p in self.items:                              # código ya declarado (F10, P1, S0…)
                refs.append({"ref": p}); continue
            m = re.match(r"^([A-Z][A-Za-z0-9\-]{0,9})\s+((?:§|cap\.|Cuadro|cuadro|tabla|Tabla|p\.|pág|secc).*)$", p)
            if m:                                            # código + localizador (CC §4.1.1)
                code = m.group(1)
                if code not in self.items:
                    self.registrar(code, f"[{code}] Sigla usada por el estudiante; completar la referencia APA.")
                refs.append({"ref": code, "loc": m.group(2)}); continue
            refs.append({"ref": self.ref_texto(p)})
        # deduplicar
        vistos, out = set(), []
        for r in refs:
            k = (r["ref"], r.get("loc"))
            if k not in vistos:
                vistos.add(k); out.append(r)
        return out


def bibliografia_del_libro(wb):
    """Precarga las listas de fuentes que algunos estudiantes declararon con código."""
    bib = Bibliografia()
    if "Instrucciones" in wb.sheetnames:                 # códigos S0, P1… en la hoja Instrucciones
        ws = wb["Instrucciones"]
        for a, b in ws.iter_rows(min_col=1, max_col=2, values_only=True):
            if a and b and re.fullmatch(r"[A-Z]{1,2}\d{1,2}", str(a).strip()):
                bib.registrar(str(a).strip(), b)
    if "Referencias" in wb.sheetnames:                   # hoja Referencias (Sigla | Referencia | Tipo | URL)
        ws = wb["Referencias"]
        for row in ws.iter_rows(min_row=4, max_col=4, values_only=True):
            if row[0] and row[1]:
                sig = re.sub(r"[^A-Za-z0-9_.\-]", "", str(row[0]))[:24]
                e = bib.registrar(sig, row[1])
                if row[3] and str(row[3]).startswith("http"):
                    bib.items[e]["url"] = str(row[3])
    for ws in wb.worksheets:                             # filas F1…F20 al pie de cada ficha
        if str(ws.cell(4, 2).value).strip() != "Clave APP":
            continue
        for r in range(5, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if a and b and re.fullmatch(r"F\d{1,2}", str(a).strip()):
                bib.registrar(str(a).strip(), b)
    return bib


# ---------------------------------------------------------------- conversión de un campo

def convertir(clave, fila, bib, avisos):
    _, _, _, uni, _, _, dato, fuente, evid, obs = fila
    d = DIC[clave]
    t, dom = d["tipo"], dominio(d["dominio"])
    c = {}
    nota = [str(obs).strip()] if obs not in (None, "") and sin_tildes(str(obs)) not in ("no disponible", "—", "-") else []
    s = "" if dato is None else str(dato).strip()
    ls = sin_tildes(s)

    if s == "":
        c["estado"] = "pendiente"
    elif ls == "no disponible" or ls.startswith("no disponible"):
        c["estado"] = "no_disponible"
    elif ls.startswith("no aplica"):
        c["estado"] = "no_aplica"
        resto = s[len("No aplica"):].strip(" .:;—-")
        if resto:
            nota.insert(0, resto)
    elif ls.startswith("pendiente"):
        c["estado"] = "pendiente"
    else:
        c["estado"] = "dato"

    if c["estado"] == "dato":
        if t in ("numero", "entero"):
            p = parse_numero(dato)
            if p:
                c.update(p)
                if isinstance(dato, str) and re.search(r"\(.*\)\s*$", dato):
                    nota.insert(0, "Texto original: " + s)
                if t == "entero" and "valor" in c:
                    c["valor"] = int(c["valor"])
                u = unidad_canonica(clave, uni)
                if u:
                    c["unidad"] = u
            else:
                c["valor_texto_original"] = s
                avisos.append(f"{clave}: número escrito como texto")
        elif t == "escala":
            try:
                n = int(float(str(dato).replace(",", ".")))
                assert 1 <= n <= 5
                c["valor"] = n
            except Exception:
                c["valor_texto_original"] = s
        elif t in ("lista", "booleano"):
            m = re.match(r"^\s*([^:(—–\-]+?)\s*(?:[:(—–\-]\s*(.*?))?\)?\s*$", s)
            base, cond = (m.group(1), m.group(2)) if m else (s, None)
            canon = en_dominio(base, dom) if dom else base
            if dom and not canon and en_dominio(s, dom):
                canon, cond = en_dominio(s, dom), None
            if canon:
                c["valor"] = canon
                if cond:
                    nota.insert(0, cond.strip())
            else:
                c["valor_texto_original"] = s
                avisos.append(f"{clave}: valor fuera de la lista controlada")
        elif t == "lista_multiple":
            sep = r"\s*\|\s*|\s*;\s*" if not dom else r"\s*\|\s*|\s*;\s*|\s*,\s*"
            items = [x.strip() for x in re.split(sep, s) if x.strip()]
            if clave == "traffic_level":
                items = [x for it in items for x in re.split(r"\s+a\s+|\s+o\s+", re.sub(r"\s*\(.*\)", "", it))]
            if clave == "imn_subregions" and sin_tildes(s).startswith("todas"):
                items = ["TODAS"]
            if dom:
                canon = [en_dominio(x, dom) for x in items]
                if all(canon):
                    c["valor"] = list(dict.fromkeys(canon))
                else:
                    c["valor_texto_original"] = s
                    avisos.append(f"{clave}: valores fuera de la lista controlada")
            else:
                c["valor"] = items
        elif t == "valor_o_clase":
            p = parse_numero(dato) if isinstance(dato, (int, float)) else None
            if p and "valor" in p:
                c["valor"] = p["valor"]
                u = unidad_canonica(clave, uni)
                if u and u != "clase":
                    c["unidad"] = u
            else:
                c["valor"] = s
        else:
            c["valor"] = s

    refs = bib.resolver(fuente)
    if refs:
        c["fuentes"] = refs
    if evid and re.fullmatch(r"E[1-5]", str(evid).strip()):
        c["evidencia"] = str(evid).strip()
    if nota:
        c["nota"] = " · ".join(nota)
    return c


# ---------------------------------------------------------------- ficha completa

def migrar_hoja(ws, bib, nuevo_id, comp, archivo):
    filas = {}
    avisos = []
    for r in range(5, ws.max_row + 1):
        fila = [ws.cell(r, col).value for col in range(1, 11)]
        k = (fila[1] or "").strip() if isinstance(fila[1], str) else None
        if not k or k not in LEGADO and k != "sources":
            continue
        filas[k] = fila
    if not filas.get("id") or filas["id"][6] in (None, ""):
        return None, ["plantilla vacía"]

    campos = {}
    for old, fila in filas.items():
        if old == "sources":
            continue
        destinos = LEGADO[old]
        if old == "weight_density":                       # se separa en densidad o masa superficial
            u = str(fila[3] or "") + " " + str(fila[6] or "")
            es_vol = re.search(r"kg/m[³3]", u) is not None
            es_sup = re.search(r"kg/m[²2](?![0-9])", u) is not None
            destinos = ["density"] if es_vol and not es_sup and "o kg" not in u else ["surface_mass"]
            if "kg/m² o kg/m³" in u:
                dv = str(fila[6] or "")
                destinos = ["density"] if re.search(r"kg/m[³3]", dv) and not re.search(r"kg/m[²2](?![0-9])", dv) else ["surface_mass"]
        for k in destinos:
            if comp not in DIC[k]["componentes"].split("|") and DIC[k]["componentes"] != "TODOS":
                continue
            nuevo = convertir(k, fila, bib, avisos)
            if k in campos:                                # dos claves legadas apuntan al mismo campo
                prev = campos[k]
                if prev["estado"] != "dato" and nuevo["estado"] == "dato":
                    campos[k] = nuevo
                elif nuevo["estado"] == "dato" and nuevo.get("valor") != prev.get("valor"):
                    prev["nota"] = (prev.get("nota", "") + f" · Valor alterno ({old}): {nuevo.get('valor', nuevo.get('valor_texto_original'))}").strip(" ·")
            else:
                campos[k] = nuevo

    campos["id"] = {"estado": "dato", "valor": nuevo_id, "fuentes": [{"ref": "PROPIO"}], "evidencia": "E1"}
    if "PROPIO" not in bib.items:
        bib.registrar("PROPIO", "Elaboración propia del estudiante (TP1).", "elaboracion_propia")
    campos["component"] = {"estado": "dato", "valor": comp}
    # campos del núcleo que la plantilla legada no tenía
    for k, d in DIC.items():
        aplica = d["componentes"] == "TODOS" or comp in d["componentes"].split("|")
        if aplica and k not in campos:
            campos[k] = {"estado": "pendiente", "nota": "Campo nuevo del estándar v1.0: completar."}
    orden = list(DIC.keys())
    campos = {k: campos[k] for k in orden if k in campos}

    usados = {r["ref"] for c in campos.values() for r in c.get("fuentes", [])}
    autor = ws.cell(2, 2).value
    ficha = {
        "schema_version": "1.0",
        "meta": {
            "id": nuevo_id, "componente": comp, "autor": str(autor).strip() if autor else "",
            "estado_revision": "borrador",
            "fecha_actualizacion": datetime.date.today().isoformat(),
            "id_legado": str(filas["id"][6]).strip(),
            "origen_migracion": {"archivo": archivo, "hoja": ws.title},
        },
        "campos": campos,
        "fuentes": [bib.items[i] for i in sorted(usados, key=lambda x: (x != "PROPIO", x))],
    }
    return ficha, avisos


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("carpeta")
    ap.add_argument("--mapa", default=str(EST / "mapa_ids_migracion.csv"))
    ap.add_argument("--salida", default=str(RAIZ / "data"))
    a = ap.parse_args()
    mapa = {(r["archivo"], r["hoja"]): r for r in csv.DictReader(open(a.mapa, encoding="utf-8"))}
    total, bitacora = 0, []
    for xl in sorted(pathlib.Path(a.carpeta).glob("*.xlsx")):
        wb = openpyxl.load_workbook(xl, data_only=True)
        bib = bibliografia_del_libro(wb)
        for ws in wb.worksheets:
            if str(ws.cell(4, 2).value).strip() != "Clave APP":
                continue
            m = mapa.get((xl.name, ws.title))
            if not m:
                bitacora.append(f"OMITIDA  {xl.name} :: {ws.title} (sin ID en el mapa)")
                continue
            ficha, avisos = migrar_hoja(ws, bib, m["id_nuevo"], m["componente"], xl.name)
            if not ficha:
                bitacora.append(f"OMITIDA  {xl.name} :: {ws.title} ({avisos[0]})")
                continue
            dest = pathlib.Path(a.salida) / CARPETA_COMP[m["componente"]] / f"{m['id_nuevo']}.json"
            dest.parent.mkdir(parents=True, exist_ok=True)
            json.dump(ficha, open(dest, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
            total += 1
            bitacora.append(f"OK       {m['id_nuevo']:12s} ← {xl.name} :: {ws.title}  ({len(avisos)} conversiones pendientes)")
    print("\n".join(bitacora))
    print(f"\n{total} fichas migradas a {a.salida}")


if __name__ == "__main__":
    main()
