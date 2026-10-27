"""Convierte una Plantilla_Ficha_SUM_<COMP>.xlsx llena a data/<componente>/<ID>.json.

Uso:  python scripts/importar_plantilla.py plantillas/mi_ficha.xlsx
"""
import csv, json, sys, re, pathlib, datetime
import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
DIC = {r["clave"]: r for r in csv.DictReader(open(RAIZ / "estandar" / "diccionario_campos.csv", encoding="utf-8"))}
LISTAS = json.load(open(RAIZ / "estandar" / "listas.json", encoding="utf-8"))
CARPETA = {c: v["nombre"].split(" ")[0].lower() for c, v in LISTAS["component"].items()}


def importar(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb["Ficha"]
    autor, fid = ws["C2"].value or "", str(ws["G2"].value or "").strip()
    comp = fid.split("-")[0]
    campos = {}
    for r in range(6, ws.max_row + 1):
        k = ws.cell(r, 2).value
        if k not in DIC:
            continue
        v = [ws.cell(r, j).value for j in range(1, 18)]
        estado, valor, vmin, vmax, cal, uni, fts, loc, evid, cert, just, nota = v[5], v[6], v[7], v[8], v[9], v[10], v[11], v[12], v[13], v[14], v[15], v[16]
        c = {"estado": estado or "pendiente"}
        t = DIC[k]["tipo"]
        if c["estado"] == "dato":
            if t == "lista_multiple" and valor is not None:
                c["valor"] = [x.strip() for x in str(valor).split("|") if x.strip()]
            elif t in ("numero", "entero", "escala") and valor is not None:
                c["valor"] = int(valor) if t != "numero" else valor
            elif valor is not None:
                c["valor"] = valor
            if vmin is not None and vmax is not None:
                c["min"], c["max"], c["calificador"] = vmin, vmax, "rango"
            if cal and "calificador" not in c:
                c["calificador"] = cal
            if uni and (t in ("numero", "entero") or t == "valor_o_clase" and isinstance(valor, (int, float))):
                c["unidad"] = uni
        if fts:
            refs = [x.strip() for x in re.split(r"[;,]", str(fts)) if x.strip()]
            c["fuentes"] = [dict(ref=x, **({"loc": str(loc)} if loc and len(refs) == 1 else {})) for x in refs]
        for clave, val in (("evidencia", evid), ("certeza", cert), ("justificacion", just)):
            if val:
                c[clave] = str(val)
        if nota and nota != DIC[k]["regla"]:
            c["nota"] = str(nota)
        campos[k] = c
    campos["id"] = {"estado": "dato", "valor": fid, "fuentes": [{"ref": "PROPIO"}], "evidencia": "E1"}
    campos["component"] = {"estado": "dato", "valor": comp}
    fuentes = []
    for row in wb["Fuentes"].iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:
            e = {"id": str(row[0]).strip(), "cita": str(row[1]).strip()}
            if row[2]: e["url"] = str(row[2])
            if row[3]: e["tipo"] = str(row[3])
            if row[4]: e["fecha_consulta"] = str(row[4])[:10]
            fuentes.append(e)
    ficha = {"schema_version": "1.0",
             "meta": {"id": fid, "componente": comp, "autor": autor, "estado_revision": "borrador",
                      "fecha_actualizacion": datetime.date.today().isoformat()},
             "campos": {k: campos[k] for k in DIC if k in campos}, "fuentes": fuentes}
    dest = RAIZ / "data" / CARPETA[comp] / f"{fid}.json"
    dest.parent.mkdir(parents=True, exist_ok=True)
    json.dump(ficha, open(dest, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"→ {dest.relative_to(RAIZ)}  (validá con: python scripts/validar.py {dest.relative_to(RAIZ)})")


if __name__ == "__main__":
    for a in sys.argv[1:]:
        importar(a)
